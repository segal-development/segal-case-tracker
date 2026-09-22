"""Regenerate the matriz reference CSVs from the business-owned xlsx.

The xlsx (``Matriz_clasificacion_procedimientos_etapas_completada.xlsx``) is
the authoritative taxonomy for the M1/M2/M3 case classification matriz. It
lives outside the repo (business-maintained, not a runtime dependency) and
must be converted once into plain CSVs committed under
``app/data/matriz/`` so the app never needs to open an xlsx at runtime.

Usage
-----
    poetry run python scripts/matriz_xlsx_to_csv.py /path/to/Matriz_*.xlsx

This overwrites:
    app/data/matriz/clasificacion.csv   (sheet "Clasificación")
    app/data/matriz/tramites.csv        (sheet "Trámites")

It does NOT regenerate ``app/data/matriz/mapeo_pjud.csv`` — that file holds
the PJUD-stage -> matriz-etapa mapping curated by hand (see
``app/services/matriz_seed.py``), which has no equivalent sheet in the xlsx.

Re-run this script whenever the business hands over an updated xlsx, review
the diff, and commit the refreshed CSVs.
"""

from __future__ import annotations

import csv
import sys
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parent.parent
OUTPUT_DIR = REPO_ROOT / "app" / "data" / "matriz"

CLASIFICACION_SHEET = "Clasificación"
TRAMITES_SHEET = "Trámites"

# Target column order for clasificacion.csv (snake_case, per the design doc).
# Maps 1:1 by position onto the "Clasificación" sheet's header row.
CLASIFICACION_COLUMNS = [
    "proc_simple_id",
    "proc_simple",
    "proc_antiguo_id",
    "proc_antiguo",
    "etapa_id",
    "etapa_padre_id",
    "etapa_padre",
    "nivel_etapa",
    "orden",
    "etapa",
    "estado_etapa",
    "condicion_rol",
    "matriz",
    "observaciones",
]

# Target column order for tramites.csv (snake_case), 1:1 by position onto the
# "Trámites" sheet's header row (10 columns).
TRAMITES_COLUMNS = [
    "id_tramite",
    "id_procedimiento",
    "procedimiento_antiguo",
    "id_etapa",
    "etapa",
    "id_etapa_padre",
    "etapa_padre",
    "nombre_del_tramite",
    "matriz_especifica_opcional",
    "observaciones",
]


def _clean_cell(value: object) -> str:
    """Normalize an xlsx cell value into the string CSV representation.

    None -> "" ; numbers -> plain str (no trailing .0 for ints) ; strings
    stripped of surrounding whitespace.
    """
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _sheet_rows(ws, expected_columns: int) -> list[list[str]]:
    """Read all data rows (skipping the header row) as cleaned string lists."""
    rows: list[list[str]] = []
    for raw_row in ws.iter_rows(min_row=2, values_only=True):
        if raw_row is None or all(v is None for v in raw_row):
            continue  # skip fully blank rows
        cleaned = [_clean_cell(v) for v in raw_row[:expected_columns]]
        rows.append(cleaned)
    return rows


def convert(xlsx_path: Path) -> None:
    if not xlsx_path.exists():
        raise FileNotFoundError(f"xlsx not found: {xlsx_path}")

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    if CLASIFICACION_SHEET not in wb.sheetnames:
        raise ValueError(f"Missing sheet {CLASIFICACION_SHEET!r} in {xlsx_path}")
    if TRAMITES_SHEET not in wb.sheetnames:
        raise ValueError(f"Missing sheet {TRAMITES_SHEET!r} in {xlsx_path}")

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    clasificacion_rows = _sheet_rows(wb[CLASIFICACION_SHEET], len(CLASIFICACION_COLUMNS))
    clasificacion_path = OUTPUT_DIR / "clasificacion.csv"
    with clasificacion_path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(CLASIFICACION_COLUMNS)
        writer.writerows(clasificacion_rows)
    print(f"Wrote {clasificacion_path} ({len(clasificacion_rows)} rows)")

    tramites_rows = _sheet_rows(wb[TRAMITES_SHEET], len(TRAMITES_COLUMNS))
    tramites_path = OUTPUT_DIR / "tramites.csv"
    with tramites_path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(TRAMITES_COLUMNS)
        writer.writerows(tramites_rows)
    print(f"Wrote {tramites_path} ({len(tramites_rows)} rows)")


def main(argv: list[str]) -> int:
    if len(argv) != 2:
        print("Usage: poetry run python scripts/matriz_xlsx_to_csv.py <path-to-xlsx>", file=sys.stderr)
        return 1
    convert(Path(argv[1]))
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
